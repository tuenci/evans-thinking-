# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_cszy_docx():
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
        
    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Georgia'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Songti SC')

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("《重生真意》")
    run_title.font.name = 'Songti SC'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # deep navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("The True Meaning of Being Born Again\n（中英双语权威典藏版）")
    run_sub.font.name = 'Georgia'
    run_sub.font.size = Pt(14)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    run_meta = p_meta.add_run("作者：王明道（Wang Mingdao，1900–1991）\n著述时间：一九三三年三月七日 · 北平（北京）")
    run_meta.font.name = 'Songti SC'
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # Lead box table
    lead_table = doc.add_table(rows=1, cols=1)
    lead_table.autofit = False
    lead_cell = lead_table.cell(0, 0)
    lead_cell.width = Inches(6.5)
    set_cell_background(lead_cell, "F7FAFC")
    lead_p = lead_cell.paragraphs[0]
    lead_p.paragraph_format.space_before = Pt(8)
    lead_p.paragraph_format.space_after = Pt(8)
    lead_p.paragraph_format.left_indent = Inches(0.15)
    lead_p.paragraph_format.right_indent = Inches(0.15)
    r_box = lead_p.add_run("全书导言：\n《重生真意》是中国教会二十世纪杰出领袖王明道先生的里程碑式属灵经典著作。全书以圣经神学第一性原理为根基，系统解明从水和圣灵生的真义、人类全然败坏的真相、新生命的活泼明证、因信称义与十字架救赎的即刻确据，以及信徒内在属灵交战中靠基督得胜有余的终极秘诀。")
    r_box.font.name = 'Songti SC'
    r_box.font.size = Pt(10)
    r_box.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_page_break()

    # Read markdown file
    with open('重生真意_王明道_中英双语对照版.md', 'r', encoding='utf-8') as f:
        md_text = f.read()

    # Split into sections by chapter
    chapters = md_text.split('\n## ')
    
    for ch_idx, ch in enumerate(chapters):
        if ch_idx == 0:
            continue # title section
            
        lines = ch.split('\n')
        ch_title_line = lines[0].strip()
        
        # Add Chapter Heading
        p_ch = doc.add_paragraph()
        p_ch.paragraph_format.space_before = Pt(20)
        p_ch.paragraph_format.space_after = Pt(12)
        r_ch = p_ch.add_run(ch_title_line)
        r_ch.font.name = 'Songti SC'
        r_ch.font.size = Pt(16)
        r_ch.font.bold = True
        r_ch.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
        # Parse bilingual pairs in this chapter
        body = '\n'.join(lines[1:])
        pairs = body.split('\n---\n')
        
        for pair in pairs:
            pair = pair.strip()
            if not pair:
                continue
            
            # Find Chinese and English
            m_zh = re.search(r'### 【中文】\s*\n(.*?)(?=\n### 【English】|\Z)', pair, re.DOTALL)
            m_en = re.search(r'### 【English】\s*\n(.*)', pair, re.DOTALL)
            
            zh_text = m_zh.group(1).strip() if m_zh else ""
            en_text = m_en.group(1).strip() if m_en else ""
            
            if zh_text:
                p_zh = doc.add_paragraph()
                p_zh.paragraph_format.space_before = Pt(6)
                p_zh.paragraph_format.space_after = Pt(4)
                p_zh.paragraph_format.line_spacing = 1.25
                r_zh_label = p_zh.add_run("【中文】 ")
                r_zh_label.font.name = 'Songti SC'
                r_zh_label.font.size = Pt(10)
                r_zh_label.font.bold = True
                r_zh_label.font.color.rgb = RGBColor(0x8C, 0x1D, 0x40) # maroon
                
                r_zh = p_zh.add_run(zh_text)
                r_zh.font.name = 'Songti SC'
                r_zh.font.size = Pt(11)
                r_zh.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
            
            if en_text:
                p_en = doc.add_paragraph()
                p_en.paragraph_format.space_before = Pt(2)
                p_en.paragraph_format.space_after = Pt(10)
                p_en.paragraph_format.line_spacing = 1.25
                r_en_label = p_en.add_run("【English】 ")
                r_en_label.font.name = 'Georgia'
                r_en_label.font.size = Pt(9.5)
                r_en_label.font.bold = True
                r_en_label.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0) # royal blue
                
                r_en = p_en.add_run(en_text)
                r_en.font.name = 'Georgia'
                r_en.font.size = Pt(10.5)
                r_en.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
                
    doc.save("重生真意_王明道_中英双语对照版.docx")
    print("Successfully generated 重生真意_王明道_中英双语对照版.docx!")

if __name__ == '__main__':
    create_cszy_docx()
